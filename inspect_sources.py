import os
import openpyxl
import pypdfium2 as pdfium
import json

DATA_DIR = r"c:\Users\samta\OneDrive\Desktop\rpl\data"

print("--- INSPECTING NQR EXCEL ---")
excel_path = os.path.join(DATA_DIR, "NQR_Qualifications_Summary.xlsx.xlsx")
wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
print(f"Sheet names: {wb.sheetnames}")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"Sheet '{sheet_name}': {len(rows)} rows")
    if rows:
        print(f"Row 0 (headers/title): {rows[0][:15]}")
        if len(rows) > 1:
            print(f"Row 1: {rows[1][:15]}")
        if len(rows) > 2:
            print(f"Row 2: {rows[2][:15]}")
        if len(rows) > 3:
            print(f"Row 3: {rows[3][:15]}")

print("\n--- INSPECTING PDFS ---")
pdf_files = [
    "Employability_Skills_30_Hours.pdf.pdf",
    "Employability_Skills_60_Hours.pdf.pdf",
    "Employability_Skills_90_Hours.pdf.pdf",
    "Employability_Skills_120_Hours.pdf.pdf",
    "Assessment_Guide.pdf.pdf"
]

for fname in pdf_files:
    fpath = os.path.join(DATA_DIR, fname)
    doc = pdfium.PdfDocument(fpath)
    n_pages = len(doc)
    sample_text = ""
    for p_idx in range(min(5, n_pages)):
        page = doc[p_idx]
        textpage = page.get_textpage()
        text = textpage.get_text_range()
        sample_text += f"\n--- Page {p_idx+1} ({len(text)} chars) ---\n" + text[:400]
    
    print(f"\nFile: {fname} | Total pages: {n_pages} | Text sample length: {len(sample_text)}")
    print(sample_text[:1200])
